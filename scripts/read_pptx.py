from pptx import Presentation
import os

path = os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop", "template.pptx")
prs = Presentation(path)
print(f"Slides: {len(prs.slides)}")
print()

for i, slide in enumerate(prs.slides):
    layout = slide.slide_layout.name if slide.slide_layout else "N/A"
    print(f"=== Slide {i+1} ({layout}) ===")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    print(f'  "{t[:150]}"')
    print()
